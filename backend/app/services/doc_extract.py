"""Text extraction for matter documents (roadmap S20).

Handles TXT / Markdown, PDF (text layer only — no OCR), and DOCX. Returns the
plain text plus a small amount of metadata; raises ExtractionError with a
user-facing message on anything it can't read.
"""
from __future__ import annotations

import io

from app.core.logging import get_logger

logger = get_logger(__name__)

TEXT_EXT = {".txt", ".md", ".markdown", ".text"}
PDF_EXT = {".pdf"}
DOCX_EXT = {".docx"}
ALLOWED_EXT = TEXT_EXT | PDF_EXT | DOCX_EXT


class ExtractionError(Exception):
    """Raised when a document cannot be turned into usable text."""


def _ext(filename: str) -> str:
    return ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""


def extract(filename: str, data: bytes) -> tuple[str, str, int | None]:
    """Return (text, media_type, page_count)."""
    ext = _ext(filename)
    if ext in TEXT_EXT:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            text = data.decode("latin-1", errors="replace")
        media = "text/markdown" if ext in {".md", ".markdown"} else "text/plain"
        return text, media, None

    if ext in PDF_EXT:
        return _extract_pdf(data)

    if ext in DOCX_EXT:
        media = (
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )
        return _extract_docx(data), media, None

    raise ExtractionError(
        f"Unsupported file type '{ext or filename}'. Upload a PDF, Word (.docx), "
        "or plain-text (.txt / .md) file."
    )


def _extract_pdf(data: bytes) -> tuple[str, str, int | None]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError("PDF support is not installed on this server.") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("The PDF could not be opened — it may be corrupt.") from exc
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ExtractionError("The PDF is password-protected.") from exc
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:  # pragma: no cover
            parts.append("")
    text = "\n\n".join(p.strip() for p in parts if p.strip())
    if not text.strip():
        raise ExtractionError(
            "No selectable text found — this looks like a scanned PDF. "
            "OCR is not supported; upload a text PDF or paste the text."
        )
    return text, "application/pdf", len(reader.pages)


def _extract_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError("Word support is not installed on this server.") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("The Word document could not be opened.") from exc
    paras = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paras.append(" | ".join(cells))
    text = "\n\n".join(paras)
    if not text.strip():
        raise ExtractionError("The Word document appears to be empty.")
    return text
